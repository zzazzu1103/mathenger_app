"""HWP + 엑셀 임포트: 문제 분리 결과와 메타데이터를 DB에 넣는다.

엑셀 '기출DB' 시트의 행 순서와 HWP 안 문제 순서가 같다고 가정하고
1번 행 ↔ 1번 문제 식으로 짝을 짓는다 (업로드된 파일 쌍으로 검증됨).
"""

from __future__ import annotations

import io
import sqlite3
from dataclasses import dataclass, field

import openpyxl

from . import db
from .hwp.reader import HwpSource
from .hwp.splitter import split_problems

# 기출DB 시트 열 → DB 필드 (0-based 열 인덱스)
EXCEL_COLUMNS = {
    1: "year", 2: "track", 3: "month", 4: "origin", 5: "number",
    6: "subject", 7: "unit_major", 8: "unit_mid", 9: "unit_small",
    10: "idea", 11: "calc_point", 12: "caution",
}

SHEET_NAME = "기출DB"


@dataclass
class ImportReport:
    source_id: int = 0
    source_name: str = ""
    n_problems: int = 0     # 문서에서 찾은 전체 문제 수
    n_meta_rows: int = 0
    n_matched: int = 0      # 메타데이터가 연결된 수
    n_added: int = 0        # 실제로 새로 등록된 수
    n_skipped: int = 0      # 중복이라 건너뛴 수
    warnings: list[str] = field(default_factory=list)


def read_excel_meta(xlsx_bytes: bytes) -> list[dict[str, str]]:
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=True, read_only=True)
    if SHEET_NAME in wb.sheetnames:
        ws = wb[SHEET_NAME]
    else:
        ws = wb.worksheets[0]
    rows: list[dict[str, str]] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row is None or len(row) == 0 or row[0] is None or str(row[0]).strip() == "":
            continue
        meta = {}
        for col, key in EXCEL_COLUMNS.items():
            value = row[col] if col < len(row) else None
            meta[key] = "" if value is None else str(value).strip()
        rows.append(meta)
    wb.close()
    return rows


def import_pair(
    conn: sqlite3.Connection,
    hwp_name: str,
    hwp_bytes: bytes,
    xlsx_bytes: bytes | None = None,
) -> ImportReport:
    """HWP(필수) + 엑셀(선택)을 임포트한다.

    이미 등록된 문제(내용이 같은 것)는 자동으로 건너뛰고 새 문제만 넣는다.
    새로 넣을 문제가 하나도 없으면 원본 파일도 저장하지 않는다.
    엑셀 없이 HWP만 올려도 되며, 메타데이터는 나중에 웹에서 수정할 수 있다.
    """
    report = ImportReport(source_name=hwp_name)

    source = HwpSource.from_bytes(hwp_bytes)
    if source.section_count() != 1:
        report.warnings.append(
            f"구역(Section)이 {source.section_count()}개인 문서입니다. 첫 구역만 사용합니다."
        )
    result = split_problems(source.body_section())
    report.warnings.extend(result.warnings)
    if not result.problems:
        raise ValueError("문서에서 문제를 하나도 찾지 못했습니다.")

    meta_rows = read_excel_meta(xlsx_bytes) if xlsx_bytes else []
    report.n_problems = len(result.problems)
    report.n_meta_rows = len(meta_rows)
    if meta_rows and len(meta_rows) != len(result.problems):
        report.warnings.append(
            f"문서에서 찾은 문제는 {len(result.problems)}개인데 엑셀 행은 {len(meta_rows)}개입니다. "
            "순서대로 짝지을 수 있는 만큼만 메타데이터를 연결했습니다."
        )

    # 이미 등록된 문제(내용 동일)와 이번 파일 안의 중복을 걸러낸다
    seen = db.existing_content_hashes(conn)
    to_add: list[tuple] = []  # (problem, meta)
    for i, problem in enumerate(result.problems):
        key = db.content_key(problem.text)
        if key in seen:
            report.n_skipped += 1
            continue
        seen.add(key)
        to_add.append((problem, meta_rows[i] if i < len(meta_rows) else {}))

    if not to_add:
        report.warnings.append(
            "모든 문제가 이미 등록되어 있어 새로 추가된 문제가 없습니다."
        )
        return report

    source_id = db.add_source(conn, hwp_name, hwp_bytes, len(to_add))
    report.source_id = source_id
    for problem, meta in to_add:
        if any(meta.values()):
            report.n_matched += 1
        db.add_problem(conn, source_id, problem.seq, problem.text, problem.blob, meta)
        report.n_added += 1
    conn.commit()
    return report
